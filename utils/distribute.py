import re
from pathlib import Path
import warnings
import logging

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd

type SheetColumnWidths = dict[str, int]

RE_INVALID_TABLE_NAME_START = re.compile(r"^[^_A-Z\\]", re.I)
RE_INVALID_TABLE_NAME_REST = re.compile(r"[^A-Z0-9_\.]", re.I)

logger = logging.getLogger(__name__)

def fill_null_columns_from_first(
    s1: pd.Series, 
    df_source: pd.DataFrame,
    match_key: str
):
    """ Replaces null fields in Series `s1` with the corresponding, non-nullish data of the matching row in DataFrame `df_source`. Matching is performed using `match_key`. If no rows match, the same Series is returned. """
    df2 = df_source[df_source[match_key] == s1.at[match_key]]

    if len(df2) > 0:
        s2: pd.Series = df2.iloc[0]

        for index, value in s1.items():
            if value == '' or pd.isna(value) or value is None:
                new_value = s2.get(index)

                if new_value != '' and not pd.isna(new_value) and new_value is not None:
                    s1.at[index] = new_value

    return s1

def get_excel_sheet_table_and_dimensions(
    excel_file: Path,
    sheet_name: str
) -> tuple[Table, SheetColumnWidths]:
    """ Reads the table and column dimension data in the specified Excel file and sheet. """
    sheet_table = None
    sheet_column_widths = None

    wb = None

    try:
        wb = load_workbook(excel_file)

        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            sheet_column_widths = get_sheet_column_widths(ws)

            ws_tables = ws.tables.items()

            if len(ws_tables) > 0:
                sheet_table = ws.tables[ws_tables[0][0]]
    except Exception as e:
        logger.debug(
            f"Failed to extract Excel table and sheet dimensions from sheet \"{sheet_name}\" of file \"{excel_file.absolute()}\".\n%s: %s",
            e.__class__.__name__,
            str(e.message) if hasattr(e, 'message') else '(No message available)'
        )
    finally:
        if wb:
            wb.close()
    
    return sheet_table, sheet_column_widths

def get_sheet_column_widths(ws: Worksheet) -> SheetColumnWidths:
    """ Reads the widths of columns in the provided worksheet according to the worksheet's reported min and max columns. Reads up to a maximum of 50 consecutive columns. """

    first_col_idx = ws.min_column
    last_col_idx = ws.max_column

    column_widths = {}

    for idx in range(first_col_idx, min(last_col_idx + 1, first_col_idx + 50), 1):
        column_letter = get_column_letter(idx)

        column_widths[column_letter] = ws.column_dimensions[column_letter].width
    
    return column_widths

def append_df_data_to_excel(
    df: pd.DataFrame,
    excel_file: Path,
    sheet_name: str = "Sheet1",
    header_row_idx: int = 0,
    drop_duplicates_indices: list | None = None,
    merge_into_existing: bool = True
) -> bool:
    """
    Appends the data in the provided DataFrame to the specified file. Assumes: 
        1. `df` and data in `excel_file` use the same format and columns
        2. First row in `excel_file` is the header row

    **NOTE:** If `merge_into_existing` is set to `False`, `drop_duplicate_indices` will have no effect.
    """
    if not merge_into_existing \
    and drop_duplicates_indices is not None:
        warnings.warn("WARNING: append_df_data_to_excel called with merge_into_existing=False and drop_duplicate_indices defined. drop_duplicate_indices will have no effect.")

    if not excel_file.is_file() or excel_file.suffix != '.xlsx':
        logger.error(f"Could not append DataFrame data. The file \"{excel_file.absolute()}\" is not a valid Excel (.xlsx) file.")

        return False
    
    if not isinstance(df, pd.DataFrame) or len(df) == 0:
        logger.warning(f"No data to append to file \"{excel_file.absolute()}\" (Sheet: {sheet_name}).")
        return False

    excel_table, sheet_column_widths = get_excel_sheet_table_and_dimensions(
        excel_file,
        sheet_name
    )
    
    # read data from the existing file
    df_existing_data = None

    if merge_into_existing:
        try:
            df_existing_data = pd.read_excel(
                excel_file,
                sheet_name=sheet_name,
                header=header_row_idx
            )
        except ValueError:
            logger.warning(f"No existing data to append. Sheet {sheet_name} does not exist in workbook \"{excel_file.absolute()}\". A new sheet will be created.")
    
    if df_existing_data is not None:
        # replace any "No Data―" instances (typically unpopulated custom fields) with empty strings
        logger.debug(
            "Read existing data from destination file. Data will be merged and duplicated will be removed. Indices for duplicate matching: %s",
            ', '.join(drop_duplicates_indices) if drop_duplicates_indices is not None and len(drop_duplicates_indices) > 0 else '(all)'
        )
        df_existing_data = df_existing_data.replace(r"(?i)No Data.*", "", regex=True)

        df_existing_data = pd.concat([df_existing_data, df], ignore_index=True)
        df_existing_data = df_existing_data.drop_duplicates(
            subset=drop_duplicates_indices,
            keep='first'
        )
        df_existing_data = df_existing_data.dropna(how='all')
    else:
        df_existing_data = df

    with pd.ExcelWriter(
        excel_file, 
        mode='a', 
        engine='openpyxl', 
        if_sheet_exists='replace'
    ) as writer:
        df_existing_data.to_excel(
            writer,
            sheet_name=sheet_name,
            index=False,
            startrow=header_row_idx
        )

    # re-create formatting and styles
    total_rows, total_cols = df_existing_data.shape

    # table begins at first row and first col
    excel_table_reference = (
        f"A{header_row_idx + 1}:"
        + str(get_column_letter(total_cols))
        + str(header_row_idx + total_rows + 1)
    )

    wb = None

    try:
        wb = load_workbook(excel_file)

        ws = wb[sheet_name]

        if sheet_column_widths is not None:
            for col, width in sheet_column_widths.items():
                ws.column_dimensions[col].width = width
            
        if excel_table is not None:
            ws.add_table(Table(
                displayName=excel_table.displayName,
                ref=excel_table_reference,
                tableStyleInfo=excel_table.tableStyleInfo
            ))
        else:
            # replace invalid table name characters with "_"
            table_name = RE_INVALID_TABLE_NAME_START.sub("_", sheet_name.strip())
            
            if len(table_name) > 1:
                table_name = table_name[0] + RE_INVALID_TABLE_NAME_REST.sub("_", table_name[1:])
            
            table_name += "_Table"

            ws.add_table(Table(
                displayName=table_name,
                ref=excel_table_reference,
                tableStyleInfo=TableStyleInfo(
                    name='TableStyleMedium2',
                    showFirstColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False
                )
            ))
        
        wb.save(excel_file)
    except Exception as e:
        logger.error(
            f"Failed to re-create Excel tables, formatting and styles for file \"{excel_file.absolute()}\".\n%s: %s",
            e.__class__.__name__,
            str(e.message) if hasattr(e, 'message') else '(No message available)'
        )
    finally:
        if wb:
            wb.close()

    return True